"""Formula-based Excel reconstruction of fitted causal dielectric models."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

def build_model_workbook(model, coefficients, wavelengths_nm, metadata=None):
    if model not in ('lorentz','drude_lorentz'):
        raise ValueError('Excel model reconstruction requires Lorentz or Drude-Lorentz coefficients')
    offset=3 if model=='drude_lorentz' else 1
    count=(len(coefficients)-offset)//3
    if count not in (1,2,3) or offset+3*count!=len(coefficients):
        raise ValueError('Invalid causal parameter count')
    wb=Workbook()
    p=wb.active;p.title='Parameters'
    p.append(['Parameter','Value','Units'])
    p.append(['epsilon_inf',float(coefficients[0]),'dimensionless'])
    if model=='drude_lorentz':
        p.append(['Drude plasma energy',float(coefficients[1]),'eV'])
        p.append(['Drude damping',float(coefficients[2]),'eV'])
    for j in range(count):
        for name,value,unit in zip((f'osc{j+1}_strength',f'osc{j+1}_energy',f'osc{j+1}_damping'),coefficients[offset+3*j:offset+3*j+3],('dimensionless','eV','eV')):
            p.append([name,float(value),unit])
    sheet=wb.create_sheet('Reconstruction')
    sheet.append(['wavelength_nm','photon_energy_eV','epsilon_real','epsilon_imag','n','k'])
    for row,w in enumerate(wavelengths_nm,2):
        sheet.cell(row,1,float(w))
        sheet.cell(row,2,f'=1239.841984/A{row}')
        real_terms=[];imag_terms=[]
        for j in range(count):
            first=offset+3*j+2 # Excel parameter row (zero-based coefficient index + 2)
            f=f'Parameters!$B${first}'
            e=f'Parameters!$B${first+1}'.replace('\\$','$')
            g=f'Parameters!$B${first+2}'.replace('\\$','$')
            a=f'({e}^2-B{row}^2)'
            b=f'({g}*B{row})'
            denom=f'({a}^2+{b}^2)'
            real_terms.append(f'({f}*{e}^2*{a}/{denom})')
            imag_terms.append(f'({f}*{e}^2*{b}/{denom})')
        if model=='drude_lorentz':
            wp='Parameters!$B$3';gd='Parameters!$B$4'
            real_terms.append(f'(-{wp}^2/(B{row}^2+{gd}^2))')
            imag_terms.append(f'({wp}^2*{gd}/(B{row}*(B{row}^2+{gd}^2)))')
        sheet.cell(row,3,'=Parameters!$B$2'+''.join('+'+t for t in real_terms))
        sheet.cell(row,4,'='+('+'.join(imag_terms) if imag_terms else '0'))
        sheet.cell(row,5,f'=SQRT((SQRT(C{row}^2+D{row}^2)+C{row})/2)')
        sheet.cell(row,6,f'=SQRT(MAX(0,(SQRT(C{row}^2+D{row}^2)-C{row})/2))')
    meta=wb.create_sheet('Metadata')
    meta.append(['Field','Value'])
    meta.append(['Model',model]);meta.append(['Oscillators',count])
    for k,v in (metadata or {}).items():meta.append([str(k),str(v)])
    notes=wb.create_sheet('Instructions')
    for line in [
        'Edit wavelength_nm values in Reconstruction; Excel recalculates epsilon, n and k.',
        'E[eV] = 1239.841984 / wavelength[nm].',
        'Lorentz epsilon = epsilon_inf + SUM(f*E0^2/(E0^2-E^2-i*gamma*E)).',
        'Drude addition = -plasma_energy^2/(E^2+i*gamma_D*E).',
        'Convention: complex refractive index n+i*k, epsilon=(n+i*k)^2.',
        'Formula results may not be cached until opened in Excel or LibreOffice.',
        'Outside the fitted spectral window, reconstruction is extrapolation, not validated measurement.'
    ]:notes.append([line])
    for ws in wb:
        ws.freeze_panes='A2'
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='17365D')
        for col in ws.columns:
            letter=get_column_letter(col[0].column)
            ws.column_dimensions[letter].width=min(88,max(16,max(len(str(c.value or '')) if c.data_type!='f' else 20 for c in col)+3))
    output=BytesIO();wb.save(output);output.seek(0);return output
